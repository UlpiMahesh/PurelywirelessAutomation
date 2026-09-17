import openpyxl
from pathlib import Path


# ============================================================
# STANDARD MARKET -> MARKET-DM MAPPING
# Source: Sheet 2 provided by user
# ============================================================

MARKET_DM_MAP = {
    "CORPUS": "CORPUS-RANA",
    "DALLAS": "DALLAS-Alimurtaza",
    "PHOENIX": "PHOENIX-Rameezuddin",
    "TAMPA": "TAMPA-DAYANE",
    "TULSA": "TULSA-ABDUL REHMAN",
    "WACO": "WACO-JUNAID",
}


# ============================================================
# HELPERS
# ============================================================

def detect_market(input_file):
    # Handles names such as:
    # Waco.xlsx, Waco1.xlsx, Waco10.xlsx
    # Phone Order Analysis - Waco(31aug).xlsx

    filename = Path(input_file).stem.upper()

    for market in sorted(MARKET_DM_MAP, key=len, reverse=True):
        if market in filename:
            return market

    raise ValueError(
        f"Could not determine the standard market from filename: "
        f"{Path(input_file).name}. "
        f"Expected one of: {', '.join(MARKET_DM_MAP.keys())}"
    )


def find_columns(ws):
    sku_col = None
    device_name_col = None
    device_count_col = None

    # Search first 5 rows because files can have different header layouts
    for row in range(1, min(6, ws.max_row + 1)):
        for col in range(1, ws.max_column + 1):

            value = ws.cell(row, col).value

            if value is None:
                continue

            header = str(value).strip().upper()

            # SKU
            if sku_col is None and header == "SKU":
                sku_col = col

            # Device name
            if device_name_col is None and header in [
                "DEVICE NAME",
                "DEVICE",
                "DESCRIPTION",
            ]:
                device_name_col = col

            # Device count
            if device_count_col is None and header in [
                "DEVICE COUNT",
                "DEVICE TOTAL",
                "COUNT",
            ]:
                device_count_col = col

    if sku_col is None:
        raise ValueError("Could not find SKU column.")

    if device_name_col is None:
        raise ValueError("Could not find Device Name column.")

    if device_count_col is None:
        raise ValueError(
            "Could not find Device Count / Device Total column."
        )

    return sku_col, device_name_col, device_count_col


# ============================================================
# STANDARD PROCESSOR
# ============================================================

def process_standard(input_file):
    input_file = Path(input_file)
    sheet_name = "OVER VIEW"

    # Detect market and MARKET-DM
    market = detect_market(input_file)
    market_dm = MARKET_DM_MAP[market]

    # Load workbook
    wb = openpyxl.load_workbook(input_file, data_only=True)

    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' was not found in {input_file.name}."
        )

    ws = wb[sheet_name]

    # Find columns
    sku_col, device_name_col, device_count_col = find_columns(ws)

    print(f"\nProcessing standard file: {input_file.name}")
    print("Market:", market)
    print("MARKET-DM:", market_dm)
    print("SKU column:", sku_col)
    print("Device Name column:", device_name_col)
    print("Device Count column:", device_count_col)

    # Extract only visible rows
    records = []

    for row in range(1, ws.max_row + 1):

        # Skip rows hidden by Excel filter/manual hiding
        if ws.row_dimensions[row].hidden:
            continue

        sku = ws.cell(row, sku_col).value
        device_name = ws.cell(row, device_name_col).value
        device_count = ws.cell(row, device_count_col).value

        # Ignore rows that aren't actual devices
        if sku is None or device_name is None:
            continue

        # Ignore headers
        if str(sku).strip().upper() == "SKU":
            continue

        records.append({
            "MARKET": market_dm,
            "SKU": sku,
            "DEVICE": device_name,
            "DEVICE COUNT": device_count,
        })

    print("Rows extracted:", len(records))

    return records


# ============================================================
# OPTIONAL STANDALONE TEST
# ============================================================

if __name__ == "__main__":
    INPUT_FILE = r"D:\ORDERS\01sep\Phone Order Analysis - Waco(31aug).xlsx"

    records = process_standard(INPUT_FILE)

    print("\nFirst 5 records:")
    for record in records[:5]:
        print(record)
