"""
Merge multiple IMEI list exports (.xls/.xlsx) into a single Excel file.

USAGE:
    python merge_imei_files.py <input_folder> <output_file.xlsx>

Example:
    python merge_imei_files.py ./uploads ./IMEI_List_Combined.xlsx

Behavior:
- Reads every .xls/.xlsx in the input folder.
- Extracts a "Market" tag from the filename (text between "IMEI_List_" and the extension).
  Falls back to the raw filename if that pattern isn't found.
- Verifies all files share the same columns before merging. If they don't,
  it stops and tells you exactly which file/columns differ instead of
  silently misaligning data.
- Outputs one combined, formatted .xlsx with a Market column, bold header,
  frozen header row, and auto-sized columns.
"""

import sys
import re
import glob
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def get_market_tag(filepath):
    name = os.path.basename(filepath)
    match = re.search(r'IMEI_List_(\w+)\.xlsx?$', name, re.IGNORECASE)
    return match.group(1) if match else os.path.splitext(name)[0]


def merge_files(input_folder, output_file):
    files = sorted(
        glob.glob(os.path.join(input_folder, '*.xls'))
        + glob.glob(os.path.join(input_folder, '*.xlsx'))
    )
    if not files:
        raise SystemExit(f'No .xls/.xlsx files found in {input_folder}')

    dfs = []
    reference_columns = None
    for f in files:
        df = pd.read_excel(f)
        if reference_columns is None:
            reference_columns = list(df.columns)
        elif list(df.columns) != reference_columns:
            raise SystemExit(
                f'Column mismatch in {os.path.basename(f)}.\n'
                f'Expected: {reference_columns}\n'
                f'Got:      {list(df.columns)}\n'
                f'Fix the source file or update the script before merging blind.'
            )
        df.insert(0, 'Market', get_market_tag(f))
        dfs.append(df)

    combined = pd.concat(dfs, ignore_index=True)
    combined.to_excel(output_file, index=False)

    wb = load_workbook(output_file)
    ws = wb.active
    for cell in ws[1]:
        cell.font = Font(name='Arial', bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Arial')
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 30)
    ws.freeze_panes = 'A2'
    wb.save(output_file)

    print(f'Merged {len(files)} files -> {output_file} ({len(combined)} rows)')
    print(combined['Market'].value_counts().to_string())


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit('Usage: python merge_imei_files.py <input_folder> <output_file.xlsx>')
    merge_files(sys.argv[1], sys.argv[2])